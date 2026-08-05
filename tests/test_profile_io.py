"""profile_io.py 동작 확인용 간단 테스트. pytest 없이 python -m tests.test_profile_io 로도 실행 가능."""

import os
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.profile_io import load_profile_from_excel  # noqa: E402

TEMPLATE = Path(__file__).resolve().parent.parent / "profile_template.xlsx"


def _make_profile_xlsx(sido: str, sigungu: str) -> str:
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["내정보"]
    ws["B3"] = 2000
    ws["C3"] = sido
    ws["D3"] = sigungu
    ws["E3"] = "미혼"
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)
    return path


def test_seoul_gu_maps_to_5digit_code():
    path = _make_profile_xlsx("서울특별시", "광진구")
    try:
        profile = load_profile_from_excel(path)
        assert profile.region_code == "11215"
    finally:
        Path(path).unlink()


def test_unknown_gu_falls_back_to_sido_code():
    path = _make_profile_xlsx("서울특별시", "테스트구")
    try:
        profile = load_profile_from_excel(path)
        assert profile.region_code == "11"
    finally:
        Path(path).unlink()


if __name__ == "__main__":
    fns = [v for k, v in list(globals().items()) if k.startswith("test_")]
    for fn in fns:
        fn()
        print(f"OK: {fn.__name__}")
    print(f"{len(fns)}개 테스트 통과")
