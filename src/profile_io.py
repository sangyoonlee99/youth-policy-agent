"""profile_template.xlsx 형식의 엑셀 파일에서 사용자 프로필을 읽어온다.

'내정보' 시트 구조:
  1행 = 헤더
  2행 = 예시 (구분 열이 "예시") — 건너뜀
  3행 = 실제 입력 (구분 열이 "내 정보")

거주지는 시/도(필수) + 시/군/구(선택)로 받는다. 시/군/구까지 매칭 테이블에 있으면
5자리 법정동코드(예: 서울 광진구 = 11215)까지 만들어서 matcher.py에 넘기고,
없으면 시/도 2자리(예: 서울 = 11)까지만 쓴다 — matcher.py의 지역 매칭은
prefix 비교라서 5자리든 2자리든 둘 다 그대로 동작한다.

주의: 이 코드표는 행정표준코드관리시스템(code.go.kr)의 법정동코드 체계를 기준으로
만들었고, 현재는 서울 25개 자치구만 들어있다. 온통청년 API의 zipCd(지역코드)
필드가 실제로 이 체계와 일치하는지는 --inspect로 응답을 받아보기 전까지 확정할 수
없다 — 다르면 이 표만 고치면 된다. 서울 외 지역은 당장은 시/도 단위까지만
매칭되고, 필요하면 SIGUNGU_CODE에 추가하면 된다.
"""

import openpyxl

from .matcher import Profile

SIDO_CODE = {
    "서울특별시": "11", "부산광역시": "26", "대구광역시": "27", "인천광역시": "28",
    "광주광역시": "29", "대전광역시": "30", "울산광역시": "31", "세종특별자치시": "36",
    "경기도": "41", "충청북도": "43", "충청남도": "44", "전라남도": "46",
    "경상북도": "47", "경상남도": "48", "제주특별자치도": "50",
    "강원특별자치도": "51", "전북특별자치도": "52",
}

# (시/도, 시/군/구) -> 5자리 법정동코드. 지금은 서울만.
SIGUNGU_CODE = {
    ("서울특별시", "종로구"): "11110", ("서울특별시", "중구"): "11140",
    ("서울특별시", "용산구"): "11170", ("서울특별시", "성동구"): "11200",
    ("서울특별시", "광진구"): "11215", ("서울특별시", "동대문구"): "11230",
    ("서울특별시", "중랑구"): "11260", ("서울특별시", "성북구"): "11290",
    ("서울특별시", "강북구"): "11305", ("서울특별시", "도봉구"): "11320",
    ("서울특별시", "노원구"): "11350", ("서울특별시", "은평구"): "11380",
    ("서울특별시", "서대문구"): "11410", ("서울특별시", "마포구"): "11440",
    ("서울특별시", "양천구"): "11470", ("서울특별시", "강서구"): "11500",
    ("서울특별시", "구로구"): "11530", ("서울특별시", "금천구"): "11545",
    ("서울특별시", "영등포구"): "11560", ("서울특별시", "동작구"): "11590",
    ("서울특별시", "관악구"): "11620", ("서울특별시", "서초구"): "11650",
    ("서울특별시", "강남구"): "11680", ("서울특별시", "송파구"): "11710",
    ("서울특별시", "강동구"): "11740",
}


_CODE_TO_GU = {code: gu for (sido, gu), code in SIGUNGU_CODE.items()}


def gu_name_from_region_code(region_code: str):
    """5자리 법정동코드에서 서울 자치구명을 역으로 찾는다. seoul_youth_client의
    sc_sarea 파라미터를 프로필 지역과 연동하기 위한 용도(하드코딩 금지 원칙).
    서울 자치구가 아니면(2자리 시/도 코드거나 매칭 안 되면) None."""
    return _CODE_TO_GU.get(str(region_code))


def load_profile_from_excel(path: str) -> Profile:
    wb = openpyxl.load_workbook(path, data_only=True)
    if "내정보" not in wb.sheetnames:
        raise ValueError(f"{path} 에 '내정보' 시트가 없습니다. profile_template.xlsx 구조를 그대로 써주세요.")
    ws = wb["내정보"]
    headers = [c.value for c in ws[1]]

    data_row = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, row))
        if record.get("구분") == "예시":
            continue
        if record.get("출생연도"):
            data_row = record
            break

    if data_row is None:
        raise ValueError(
            f"{path} 의 '내정보' 시트에서 실제 입력 행을 찾지 못했습니다. "
            f"3행(노란색)의 '출생연도'부터 채워주세요."
        )

    sido = str(data_row.get("거주_시도") or "").strip()
    sigungu = str(data_row.get("거주_시군구") or "").strip()

    region_code = SIDO_CODE.get(sido, "")
    if sido and not region_code:
        raise ValueError(f"'{sido}' 는 알 수 없는 시/도명입니다. '지역코드표' 시트의 표기를 그대로 써주세요.")

    if sido and sigungu:
        fine_code = SIGUNGU_CODE.get((sido, sigungu))
        if fine_code:
            region_code = fine_code
        else:
            print(
                f"[안내] '{sido} {sigungu}' 는 아직 구/군 단위 코드표에 없어서 "
                f"'{sido}' 단위(시/도)까지만 지역 매칭에 씁니다. 더 정밀하게 하려면 "
                f"src/profile_io.py의 SIGUNGU_CODE에 추가하세요."
            )

    interests_raw = data_row.get("관심분야") or ""
    interests = [s.strip() for s in str(interests_raw).split(",") if s.strip()]

    return Profile(
        birth_year=int(data_row["출생연도"]),
        region_code=region_code,
        marital_status=data_row.get("혼인여부") or "무관",
        annual_income_10k=data_row.get("연소득_만원"),
        interests=interests,
        # 아래 4개는 온통청년 코드정의서 라벨과 정확히 맞춘 드롭다운 값 — 하드
        # 필터에도 쓰인다(matcher.py의 CODE_MAPS/_CODE_FILTER_FIELDS 참고).
        education_status=str(data_row.get("학적상태") or ""),
        employment_status=str(data_row.get("취업상태") or ""),
        specialization=str(data_row.get("전공분야") or ""),
        household_type=str(data_row.get("가구유형") or ""),
        # 아래는 하드 필터엔 안 쓰고 LLM 소프트 필터 참고용 자유 텍스트.
        employment_detail=str(data_row.get("취업상태_상세") or ""),
        specialization_detail=str(data_row.get("전공상세") or ""),
        military_status=str(data_row.get("병역상태") or ""),
        notes=str(data_row.get("기타메모") or ""),
    )
